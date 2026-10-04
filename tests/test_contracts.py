import copy,io,json,sys,tempfile,unittest,sqlite3
from pathlib import Path
from decimal import Decimal
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from plfss_service.model import amount,branch,stage_for
from plfss_service.store import Store
from plfss_service.web import filters,export
from tools.normalize import table_facts
from tools.audit_source import check
import lxml.html
from openpyxl import load_workbook

class Contracts(unittest.TestCase):
 def setUp(self):self.store=Store(ROOT/'data')
 def test_blank_is_not_zero(self):
  for value in ('',' ','-','—','n.d.'):self.assertIsNone(amount(value,10**9))
  self.assertEqual(amount('0,0',10**9),0)
 def test_unit_negative_and_spacing(self):
  self.assertEqual(amount('− 12,7',10**9),-1270000000000)
  self.assertEqual(amount('1\u202f234,56',1),123456)
 def test_phases(self):
  self.assertEqual(stage_for('PLFSS',2027,2027),'PLFSS')
  self.assertEqual(stage_for('PLFSS',2027,2026),'PLFSS_RECTIF')
  self.assertEqual(stage_for('PLFSS',2027,2025),'CONSTATE')
  self.assertEqual(stage_for('LFSS',2026,2027),'PROJECTION')
 def test_entities_with_fsv(self):
  self.assertEqual(branch('Régimes obligatoires de base de sécurité sociale consolidés'),'TOTAL')
  self.assertEqual(branch('Toutes branches y compris Fonds de solidarité vieillesse'),'TOTAL_FSV')
  self.assertEqual(branch('Fonds de solidarité vieillesse'),'FSV')
 def test_2027_initial_values(self):
  for metric,val in [('RECETTES',68370000000000),('DEPENSES',69640000000000),('SOLDE',-1270000000000)]:
   m=self.store.matrix('EQUILIBRE','ROBSS',2027,2027,metric,('PLFSS','LFSS','CONSTATE'))
   self.assertEqual(m['total']['cells'][0]['amount_cents'],val)
   self.assertIsNone(m['total']['cells'][1]['amount_cents']);self.assertIsNone(m['total']['cells'][2]['amount_cents'])
 def test_2026_voted_is_not_proposed(self):
  m=self.store.matrix('EQUILIBRE','ROBSS',2026,2026,'DEPENSES',('PLFSS','LFSS'))
  self.assertEqual([c['amount_cents'] for c in m['total']['cells']],[67690000000000,68420000000000])
 def test_2019_rg_not_all_regimes(self):
  a=self.store.matrix('EQUILIBRE','RG',2019,2019,'RECETTES',('LFSS',))
  b=self.store.matrix('EQUILIBRE','ROBSS',2019,2019,'RECETTES',('LFSS',))
  self.assertEqual(a['total']['cells'][0]['amount_cents'],40480000000000)
  self.assertEqual(b['total']['cells'][0]['amount_cents'],51090000000000)
 def test_2017_2018_direct_voted_tables(self):
  for year in (2017,2018):
   m=self.store.matrix('EQUILIBRE','ROBSS',year,year,'DEPENSES',('PLFSS','LFSS'))
   for c in m['total']['cells']:self.assertEqual(c['status'],'published')
 def test_ondam_not_added_to_branches(self):
  a=self.store.matrix('ONDAM','ROBSS',2027,2027,'DEPENSES',('PLFSS',))
  self.assertEqual(a['total']['cells'][0]['amount_cents'],27870000000000)
  self.assertNotIn('MALADIE',[r['entity'] for r in a['rows']])
 def test_historical_ondam_actual_columns(self):
  for year,expected in ((2017,19070000000000),(2018,19510000000000),(2019,20030000000000)):
   m=self.store.matrix('ONDAM','ROBSS',year,year,'DEPENSES',('CONSTATE',))
   self.assertEqual(m['total']['cells'][0]['amount_cents'],expected)
   for r in [m['total'],*m['rows']]:
    self.assertEqual(r['cells'][0]['status'],'published')
    proof=self.store.proof(r['cells'][0]['references'][0])
    self.assertEqual(proof['column_label'],'Constat '+str(year))
    self.assertIsInstance(proof['page_number'],int)
 def test_pdf_auditor_detects_wrong_actual_column_and_page(self):
  m=self.store.matrix('ONDAM','ROBSS',2019,2019,'DEPENSES',('CONSTATE',))
  fact=self.store.proof(m['total']['cells'][0]['references'][0]);source=fact['source']
  self.assertIsNone(check(fact,source,ROOT/'data'))
  for field,value in [('amount_cents',fact['amount_cents']+1),('column_index',2),('exercise',2020),('unit_eur',10**6),('pdf_bbox',[0,0,1,1])]:
   changed=copy.deepcopy(fact);changed[field]=value
   self.assertIsNotNone(check(changed,source,ROOT/'data'),field)
 def test_excel_auditor_reads_year_and_unit_from_source(self):
  f=json.loads((ROOT/'data/extracted/excel-facts.json').read_text(encoding='utf-8'))[0]
  source=self.store.source(f['source_id']);self.assertIsNone(check(f,source,ROOT/'data'))
  for field,value in [('exercise',2017),('number_format','General'),('header_cell','A1'),('label_cell','A1')]:
   changed=copy.deepcopy(f);changed[field]=value
   self.assertIsNotNone(check(changed,source,ROOT/'data'),field)
 def test_abbreviated_actual_year_is_not_revision_date(self):
  facts=json.loads((ROOT/'data/extracted/excel-facts.json').read_text(encoding='utf-8'))
  for sid,year in [('edb248c41c2cb0ef94a50a92',2022),('bb6eb7d889e0cef8c3d6d2e6',2021)]:
   found=[f for f in facts if f['source_id']==sid and f['header_cell']=='D3']
   self.assertEqual(len(found),7)
   self.assertTrue(all(f['exercise']==year for f in found))
 def test_fsv_only_does_not_include_combined_regime_accounts(self):
  m=self.store.matrix('EQUILIBRE','FSV',2020,2020,'DEPENSES',('LFSS',))
  self.assertEqual(m['total']['cells'][0]['amount_cents'],1820000000000)
  combined=self.store.matrix('EQUILIBRE','RG_FSV',2020,2020,'DEPENSES',('LFSS',))
  self.assertEqual(combined['total']['cells'][0]['amount_cents'],41510000000000)
 def test_latest_projection_retains_its_legislative_snapshot(self):
  candidates=[dict(exercise=2027,stage='PROJECTION',entity='TOTAL',priority=50,edition=2026,kind=kind,id=i,amount_cents=value,precision_eur='0.01',raw=str(value/100),unit_eur=1,source_id=str(i)) for i,kind,value in [(1,'PLFSS',10000),(2,'LFSS',11000)]]
  cell=self.store.cell(candidates,2027,'PROJECTION','TOTAL')
  self.assertEqual(cell['amount_cents'],11000)
  self.assertEqual(cell['earlier_references'],[1])
 def test_merged_html_annex_keeps_actual_column_and_entity(self):
  with self.store.connect() as db:
   facts=[dict(json.loads(r[0])) for r in db.execute("SELECT payload FROM facts WHERE source_id='d5a97207cd2a12e2bd0c4177' AND exercise=2017 AND metric='DEPENSES' AND priority=50")]
  rg=next(f for f in facts if f['perimeter']=='RG' and f['entity']=='VIEILLESSE')
  self.assertEqual(rg['raw'],'124,9');self.assertEqual(rg['column_index'],4)
  self.assertIsNone(check(rg,self.store.source(rg['source_id']),ROOT/'data'))
  for field,value in [('entity','TOTAL'),('exercise',2018),('metric','RECETTES'),('year_header_column_index',6)]:
   changed=copy.deepcopy(rg);changed[field]=value;self.assertIsNotNone(check(changed,self.store.source(rg['source_id']),ROOT/'data'),field)
 def test_source_notice_preserves_value_and_follows_proofs(self):
  m=self.store.matrix('EQUILIBRE','RG_FSV',2018,2018,'SOLDE',('CONSTATE',))
  cell=m['total']['cells'][0]
  self.assertEqual(cell['amount_cents'],120000000000)
  self.assertTrue(cell['source_notices'])
  self.assertIn('Rapprochement',cell['reason'])
  self.assertTrue(self.store.proof(cell['references'][0])['source_notices'])
 def test_source_notice_is_invalidated_by_changed_value(self):
  with self.store.connect() as db:
   f=json.loads(db.execute("SELECT payload FROM facts WHERE source_id='JORFARTI000039675325' AND perimeter='RG_FSV' AND metric='SOLDE'").fetchone()[0])
  self.assertTrue(self.store.notices(f));f['amount_cents']+=1
  self.assertFalse(self.store.notices(f))
 def test_all_published_matrices_have_no_unexplained_conflict(self):
  for p in ('RG','ROBSS','RG_FSV','ROBSS_FSV','FSV'):
   for metric in ('RECETTES','DEPENSES','SOLDE'):
    m=self.store.matrix('EQUILIBRE',p,2017,2027,metric,('PLFSS','LFSS','CONSTATE','PLFSS_RECTIF','LFSS_RECTIF','PROJECTION'))
    self.assertFalse(any(c['status']=='divergence' for row in [m['total'],*m['rows']] for c in row['cells']),(p,metric))
 def test_exclusions_preserve_consolidated_total(self):
  normal=self.store.matrix('EQUILIBRE','ROBSS',2027,2027,'DEPENSES',('PLFSS',))
  m=self.store.matrix('EQUILIBRE','ROBSS',2027,2027,'DEPENSES',('PLFSS',),('MALADIE',))
  self.assertEqual(normal['total'],m['total'])
  self.assertEqual(m['selection']['cells'][0]['amount_cents'],sum(r['cells'][0]['amount_cents'] for r in m['rows'] if r['entity']!='MALADIE'))
  self.assertIn('transferts non neutralisés',m['selection']['label'])
 def test_empty_selection_remains_distinct_from_zero(self):
  m=self.store.matrix('ONDAM','ROBSS',2027,2027,'DEPENSES',('PLFSS',),('SOINS_VILLE','ETABLISSEMENTS_SANTE','PERSONNES_AGEES','PERSONNES_HANDICAPEES','FIR','AUTRES'))
  self.assertEqual(m['selection']['cells'][0]['status'],'empty_selection')
  self.assertIsNone(m['selection']['cells'][0]['amount_cents'])
  self.assertEqual(m['total']['cells'][0]['amount_cents'],27870000000000)
 def test_fsv_has_no_duplicate_branch_row(self):
  m=self.store.matrix('EQUILIBRE','FSV',2017,2027,'DEPENSES',('PLFSS','LFSS','CONSTATE'))
  self.assertEqual(m['rows'],[])
  self.assertEqual(m['total']['entity'],'FSV')
 def test_reject_invalid_filters(self):
  for p in ({'start':['2027'],'end':['2017']},{'domain':['ONDAM'],'metric':['SOLDE']},{'perimeter':['../../data']},{'exclude':['SOINS_VILLE']},{'stages':['PLFSS,PLFSS']},{'perimeter':['FSV'],'exclude':['MALADIE']}):
   with self.assertRaises(ValueError):filters(p)
 def test_document_year_filter(self):
  for year in (2017,2026,2027):
   docs=self.store.documents(year)
   self.assertTrue(docs);self.assertTrue(all(d['publication_year']==year for d in docs))
 def test_numeric_excel_display(self):
  m=self.store.matrix('EQUILIBRE','ROBSS',2027,2027,'DEPENSES',('PLFSS','LFSS'))
  body,mime=export(m,'xlsx');book=load_workbook(io.BytesIO(body));ws=book['Comptes sociaux']
  self.assertEqual(ws['F2'].value,696400000000);self.assertEqual(ws['F2'].data_type,'n');self.assertEqual(ws['F2'].number_format,'#,##0.00;[Red]-#,##0.00;0.00');self.assertIsNone(ws['F3'].value)
 def test_french_fragment_regression(self):
  html='<p>Pour l’année 2027, tableau de l’ensemble des régimes obligatoires de base.</p><p>(En milliards d’euros)</p><table><tr><th></th><th>Recettes</th><th>Dépenses</th><th>Solde</th></tr><tr><td>Maladie</td><td>1,2</td><td>1,3</td><td>− 0,1</td></tr></table>'
  f,_=table_facts(lxml.html.fromstring(html),{'id':'test','sha256':'test'},2027,'PLFSS')
  self.assertEqual(len(f),3);self.assertTrue(all(x['exercise']==2027 for x in f));self.assertEqual(f[1]['metric'],'DEPENSES')
 def test_source_auditor_detects_changed_amount_unit_cell_hash(self):
  with self.store.connect() as db:
   f=json.loads(db.execute("SELECT payload FROM facts WHERE exercise=2027 AND stage='PLFSS' AND entity='TOTAL' AND metric='DEPENSES' AND priority=100").fetchone()[0])
  source=self.store.source(f['source_id']);self.assertIsNone(check(f,source,ROOT/'data'))
  for field,value in [('amount_cents',f['amount_cents']+1),('unit_eur',1000000),('column_index',1),('source_sha','0'*64),('exercise',2018)]:
   changed=copy.deepcopy(f);changed[field]=value;self.assertIsNotNone(check(changed,source,ROOT/'data'),field)
 def test_current_matrix_has_no_unexplained_conflict(self):
  for metric in ('RECETTES','DEPENSES','SOLDE'):
   m=self.store.matrix('EQUILIBRE','ROBSS',2017,2027,metric,('PLFSS','LFSS','CONSTATE'))
   self.assertTrue(all(c['status']!='divergence' for row in [m['total'],*m['rows']] for c in row['cells']))

if __name__=='__main__':unittest.main()
