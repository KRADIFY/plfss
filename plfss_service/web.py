"""Separate public PLFSS web service. No scheduler, paid service or production-site dependency."""
import csv,io,json,logging,mimetypes,os,re,traceback
from decimal import Decimal
from http.server import BaseHTTPRequestHandler,ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlsplit,parse_qs,unquote
from .store import Store,PERIMETERS
from .model import STAGES,METRICS
from . import retrieval_client
ROOT=Path(__file__).resolve().parents[1]
DATA=Path(os.getenv('PLFSS_DATA',str(ROOT/'data')))
STORE=Store(DATA)
PUBLIC=ROOT/'public'

def number(params,key,default,lo,hi):
 value=params.get(key,[str(default)])[0]
 if not re.fullmatch('[0-9]{1,4}',value):raise ValueError('Nombre invalide : '+key)
 value=int(value)
 if not lo<=value<=hi:raise ValueError('Valeur hors périmètre : '+key)
 return value

def filters(params):
 domain=params.get('domain',['EQUILIBRE'])[0];perimeter=params.get('perimeter',['ROBSS'])[0]
 if domain not in ('EQUILIBRE','ONDAM') or perimeter not in PERIMETERS:raise ValueError('Périmètre inconnu')
 if domain=='ONDAM' and perimeter!='ROBSS':raise ValueError('L’ONDAM a son propre périmètre ; choisir Tous les régimes de base.')
 metric=params.get('metric',['DEPENSES'])[0]
 if metric not in METRICS or (domain=='ONDAM' and metric!='DEPENSES'):raise ValueError('Indicateur invalide')
 start=number(params,'start',2023,2017,2027);end=number(params,'end',2027,2017,2027)
 if start>end:raise ValueError('L’année de début doit précéder celle de fin.')
 stages=tuple(params.get('stages',['PLFSS,LFSS,CONSTATE'])[0].split(','))
 if not stages or len(set(stages))!=len(stages) or any(x not in STAGES for x in stages):raise ValueError('Étape inconnue ou répétée')
 allowed=({'SOINS_VILLE','ETABLISSEMENTS_SANTE','PERSONNES_AGEES','PERSONNES_HANDICAPEES','FIR','AUTRES'} if domain=='ONDAM' else {'MALADIE','ATMP','VIEILLESSE','FAMILLE','AUTONOMIE'}) if perimeter in ('ROBSS','RG') else set()
 excluded=tuple(x for x in params.get('exclude',[''])[0].split(',') if x)
 if any(x not in allowed for x in excluded):raise ValueError('Poste inconnu')
 return dict(domain=domain,perimeter=perimeter,start=start,end=end,metric=metric,stages=stages,excluded=excluded)

def records(matrix):
 rows=[matrix['total']]+([matrix['selection']] if matrix['selection'] else [])+matrix['rows'];out=[]
 for row in rows:
  for col,cell in zip(matrix['columns'],row['cells']):
   out.append([col['year'],STAGES[col['stage']],PERIMETERS[matrix['perimeter']],row['label'],METRICS[matrix['metric']],
     None if cell['amount_cents'] is None else Decimal(cell['amount_cents'])/100,cell['status'],cell.get('reason',''),','.join(map(str,cell['references'])),row.get('excluded',False)])
 return out

def export(matrix,kind):
 headers=['Exercice','Étape','Périmètre','Poste','Indicateur','Montant EUR','Statut','Explication','Preuves','Poste exclu']
 if kind=='csv':
  out=io.StringIO();writer=csv.writer(out,delimiter=';');writer.writerow(headers)
  for row in records(matrix):writer.writerow([("'"+str(v)) if isinstance(v,str) and v.startswith(('=','+','-','@')) else v for v in row])
  return out.getvalue().encode('utf-8-sig'),'text/csv; charset=utf-8'
 if kind!='xlsx':raise ValueError('Format inconnu')
 from openpyxl import Workbook
 from openpyxl.styles import Font,PatternFill,Alignment
 from openpyxl.utils import get_column_letter
 wb=Workbook();ws=wb.active;ws.title='Comptes sociaux';ws.append(headers)
 for row in records(matrix):
  ws.append(row)
  for c in ws[ws.max_row]:
   if isinstance(c.value,str):c.data_type='s'
 for c in ws[1]:c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='173A55')
 for row in ws.iter_rows(min_row=2):
  row[5].number_format='#,##0.00;[Red]-#,##0.00;0.00';row[7].alignment=Alignment(wrap_text=True,vertical='top')
 for i,w in enumerate((12,28,38,65,16,28,18,75,28,16),1):ws.column_dimensions[get_column_letter(i)].width=w
 ws.freeze_panes='F2';ws.auto_filter.ref=ws.dimensions
 memo=wb.create_sheet('Lire les chiffres');memo.append(['Règle','Explication']);memo.append(['Périmètre',matrix['note']]);memo.append(['Zéros','Un zéro publié reste 0 ; une absence reste une cellule vide, expliquée par le statut.']);memo.append(['Précision','Les nombres restent numériques. Le format Excel fixe évite la notation exponentielle ; les sources budgétaires peuvent être arrondies.']);memo.append(['Version',STORE.meta()['data_version']]);memo.column_dimensions['A'].width=20;memo.column_dimensions['B'].width=110
 for r in memo.iter_rows():
  for c in r:c.alignment=Alignment(wrap_text=True,vertical='top')
 stream=io.BytesIO();wb.save(stream);return stream.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

class Handler(BaseHTTPRequestHandler):
 def log_message(self,fmt,*args):logging.info(fmt,*args)
 def send(self,status,body,mime='application/json; charset=utf-8',filename=None):
  if isinstance(body,(dict,list)):body=json.dumps(body,ensure_ascii=False).encode('utf-8')
  if isinstance(body,str):body=body.encode('utf-8')
  self.send_response(status);self.send_header('Content-Type',mime);self.send_header('Content-Length',str(len(body)))
  self.send_header('X-Content-Type-Options','nosniff');self.send_header('Referrer-Policy','same-origin')
  self.send_header('Content-Security-Policy',"default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self'; connect-src 'self'; font-src 'self'; object-src 'none'; frame-ancestors 'none'; base-uri 'none'")
  self.send_header('Cache-Control','no-store' if mime.startswith('application/json') else 'no-cache')
  if filename:self.send_header('Content-Disposition','attachment; filename="'+filename+'"')
  self.end_headers();self.wfile.write(body)
 def file(self,path,mime=None):
  if not path.is_file():self.send(404,dict(error='Fichier absent du lot local.'));return
  size=path.stat().st_size;start=0;end=size-1;status=200
  range_header=self.headers.get('Range')
  if range_header:
   match=re.fullmatch(r'bytes=(\d+)-(\d*)',range_header)
   if not match:self.send(416,dict(error='Plage non prise en charge'));return
   start=int(match[1]);end=min(size-1,int(match[2])) if match[2] else size-1
   if start>=size or end<start:self.send(416,dict(error='Plage hors fichier'));return
   status=206
  self.send_response(status);self.send_header('Content-Type',mime or mimetypes.guess_type(path.name)[0] or 'application/octet-stream');self.send_header('Content-Length',str(end-start+1));self.send_header('Accept-Ranges','bytes');self.send_header('X-Content-Type-Options','nosniff')
  if status==206:self.send_header('Content-Range',f'bytes {start}-{end}/{size}')
  self.end_headers()
  with path.open('rb') as stream:
   stream.seek(start);remaining=end-start+1
   while remaining:
    buf=stream.read(min(65536,remaining))
    if not buf:break
    self.wfile.write(buf);remaining-=len(buf)
 def do_GET(self):
  try:
   parsed=urlsplit(self.path);path=unquote(parsed.path);params=parse_qs(parsed.query)
   if path=='/healthz':self.send(200,dict(ready=STORE.data.joinpath('derived/plfss.sqlite').is_file()));return
   if path=='/api/meta':self.send(200,STORE.meta());return
   if path=='/api/search/status':self.send(200,retrieval_client.request('/healthz'));return
   if path=='/api/search':self.send(200,retrieval_client.resolve_sources(retrieval_client.search(params),STORE));return
   if path.startswith('/api/passage/'):
    self.send(200,retrieval_client.resolve_sources(retrieval_client.passage(path.rsplit('/',1)[-1]),STORE));return
   if path=='/api/matrix':self.send(200,STORE.matrix(**filters(params)));return
   if path=='/api/documents':
    year=number(params,'year',2027,2017,2027) if params.get('year',[''])[0] else None
    docs=STORE.documents(year,params.get('family',[''])[0],params.get('q',[''])[0][:250])
    offset=number(params,'offset',0,0,9999)
    self.send(200,dict(total=len(docs),offset=offset,items=docs[offset:offset+100],filter_meaning='Millésime du document ; les années des montants sont traitées séparément.'));return
   if path=='/api/proof':
    f=STORE.proof(number(params,'id',0,1,9999999));self.send(200 if f else 404,f or dict(error='Preuve non trouvée'));return
   if path=='/api/coverage':
    target=ROOT/'reports/coverage.json'
    if not target.is_file():self.send(200,dict(status='Bilan documentaire en préparation',rows=[]));return
    self.send(200,json.loads(target.read_text(encoding='utf-8')));return
   if path=='/api/export':
    payload,mime=export(STORE.matrix(**filters(params)),params.get('format',['xlsx'])[0]);self.send(200,payload,mime,'PLFSS-comptes-sociaux.'+params.get('format',['xlsx'])[0]);return
   if path.startswith('/documents/'):
    sid=path.removeprefix('/documents/');source=STORE.source(sid)
    if not source or source.get('status')!='downloaded':self.send(404,dict(error='Document non récupéré.'));return
    target=(DATA/source['path']).resolve()
    if not target.is_relative_to(DATA.resolve()):raise ValueError('Chemin invalide')
    # Untrusted source HTML is never executed on the application origin.
    if target.suffix.lower() in ('.html','.htm','.asp'):
     self.send(400,dict(error='Utilisez la source officielle ou la fenêtre de preuve pour les documents HTML.'));return
    self.file(target);return
   target=(PUBLIC/('index.html' if path=='/' else path.lstrip('/'))).resolve()
   if not target.is_relative_to(PUBLIC.resolve()):raise ValueError('Chemin invalide')
   self.file(target)
  except (ValueError,KeyError) as exc:self.send(400,dict(error=str(exc)))
  except (BrokenPipeError,ConnectionResetError):pass
  except Exception:
   logging.exception('PLFSS request failed');self.send(500,dict(error='Le service n’a pas pu terminer cette demande. Réessayez.'))

def main():
 logging.basicConfig(level=logging.INFO)
 server=ThreadingHTTPServer(('0.0.0.0',int(os.getenv('PORT','18895'))),Handler);server.daemon_threads=True;server.serve_forever()
if __name__=='__main__':main()
